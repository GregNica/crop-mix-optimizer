"""
build_toy_v1.py — generates the Spine + Storage toy model workbook.

WHY A GENERATOR AND NOT A HAND-BUILT SHEET. The calendar (which week each succession
plants, grows and harvests) and every crop constant are read from the SAME source the app
uses: farm_v2.utils.calendar.generate_successions and tests/golden/default_farm_payload.json.
A hand-typed sheet would be a second guess at the model, and the whole point of this toy is
to be an INDEPENDENT CHECK on the app's arithmetic — which it cannot be if its inputs are
a retyped approximation. Re-run this after any crop-data change.

Output: farm_toy_v1_spine_storage.xlsx  (open in LibreOffice Calc or Excel)
"""
import json, sys, pathlib

REPO = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

from farm_v2.utils.calendar import generate_successions, harvest_weeks_of

# ---------------------------------------------------------------------------
# Source data — from the app, not retyped
# ---------------------------------------------------------------------------
payload = json.load(open(REPO / "tests" / "golden" / "default_farm_payload.json"))
CROP_NAMES = ["Spinach", "Kale", "Garlic"]
crops = {c["name"]: c for c in payload["crops"] if c["name"] in CROP_NAMES}

# Five successions, chosen to show: a spring and a fall planting of each quick crop, and the
# one overwintering crop. Picked from the 68 the real generator offers for these three crops.
PICKED = ["Spinach W10", "Spinach W34", "Kale W12", "Kale W32", "Garlic W47"]
all_succ = {s["name"]: s for s in generate_successions(list(crops.values()), 2024)}
succs = [all_succ[n] for n in PICKED]

STF        = payload["sell_through_fraction"]      # 0.85
COLD_CAP   = payload["harvest_storage_cap"]        # 275 m3
AMB_CAP    = payload["ambient_storage_cap"]        # 150 m3
TOTAL_LAND = payload["total_land"]                 # 80 acres

W0 = 3                       # week 1 lives in column C
def wcol(w): return W0 + w - 1
def wl(w):   return get_column_letter(wcol(w))
LAST = get_column_letter(wcol(52))

# ---------------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------------
F_TITLE  = Font(bold=True, size=15)
F_H1     = Font(bold=True, size=12, color="1F4E79")
F_BOLD   = Font(bold=True)
F_NOTE   = Font(italic=True, size=9, color="666666")
F_SMALL  = Font(size=9)
IN_FILL  = PatternFill("solid", fgColor="FFF2CC")   # yellow = you type here
CALC_FILL= PatternFill("solid", fgColor="F2F2F2")
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
POOL_C   = PatternFill("solid", fgColor="E8F0FE")
POOL_A   = PatternFill("solid", fgColor="FFF0E6")
THIN     = Side(style="thin", color="BFBFBF")
BOX      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

# ===========================================================================
# SHEET 1 — READ ME
# ===========================================================================
rd = wb.active
rd.title = "READ ME"
rd.column_dimensions["A"].width = 112
readme = [
 ("Farm toy model — v1: THE SPINE + STORAGE", F_TITLE),
 ("", None),
 ("WHAT THIS IS", F_H1),
 ("A hand-checkable model of the smallest part of the app that you did not build in Excel:", None),
 ("what you plant becomes what you harvest becomes what you can hold becomes what you sell.", None),
 ("There is NO solver in this version. You type the acres. Everything else is arithmetic you", None),
 ("can follow cell by cell.", None),
 ("", None),
 ("Three crops, chosen to make one distinction visible: Spinach and Kale are refrigerated", None),
 ("(the cooler); Garlic is cured and lives in the barn at ambient temperature. They compete", None),
 ("for two SEPARATE capped spaces, which is why the app has two storage pools.", None),
 ("", None),
 ("HOW TO USE IT", F_H1),
 ("1. Go to the 'Model' sheet. Every YELLOW cell is an input. Nothing else should be typed in.", None),
 ("2. Start with the acres in B6:B10. Change one. Watch the whole sheet move.", None),
 ("3. Watch row 'COLD POOL headroom' and 'AMBIENT POOL headroom'. Red = you have overflowed", None),
 ("   that store in that week. The app would call this plan infeasible.", None),
 ("", None),
 ("THE ONE THING THIS SHEET EXISTS TO SHOW", F_H1),
 ("Set every crop's 'Sell rate lb/wk' high enough that stock never builds up. Closing inventory", None),
 ("goes to zero in every week — you are selling everything the week you pick it.", None),
 ("Now look at the cold pool. IT IS STILL FULL IN THE HARVEST WEEKS.", F_BOLD),
 ("", None),
 ("That is the whole point of the storage constraint, and it is the thing the app got wrong", None),
 ("for months. A cooler is sized by what passes THROUGH it, not by what sleeps in it. Field", None),
 ("heat has to come off spinach within the hour whether it ships Tuesday or Friday. The old", None),
 ("version measured end-of-week stock, so a farm that sold everything immediately was reported", None),
 ("as needing NO cooler at all.", None),
 ("", None),
 ("The formula that encodes this is on the 'Room occupancy' rows:", None),
 ("      room this week = (stock you carried in) + (everything you harvest this week)", F_BOLD),
 ("...counted BEFORE any of it ships. You size a cooler for the worst Tuesday, not the average.", None),
 ("", None),
 ("WHAT THIS VERSION DOES NOT CARRY", F_H1),
 ("Stated plainly so you know the edges, and so nothing here is a surprise later:", None),
 ("  - No demand caps. 'Sell rate lb/wk' is a stand-in for the buyer. The real model has", None),
 ("    per-channel weekly AND annual caps, which is layer 2.", None),
 ("  - No prices, costs, revenue or profit. This layer is physical only.", None),
 ("  - No cash, no labour, no water, no crop diversity, no planting/harvest throughput limits.", None),
 ("  - Shelf life is shown as a WARNING, not enforced, and it ignores the decay weighting the", None),
 ("    real constraint applies. See the 'Over shelf life?' rows.", None),
 ("  - The real model DECIDES the acres. Here you type them. That is layer 2 as well.", None),
 ("", None),
 ("WHY 'Sold' USES MIN() - AND WHY THAT IS A FINDING, NOT A SHORTCUT", F_H1),
 ("The 'Sold' row is MIN(what you have, what you can move). That is fine here because nothing", None),
 ("is being optimised. But MIN() is NOT LINEAR, so a linear solver cannot use it. This is", None),
 ("exactly why the real model does not have a 'sold' formula at all: it makes the pounds sold a", None),
 ("DECISION VARIABLE and bounds it with a constraint instead. When the spreadsheet cannot", None),
 ("express something, that is usually telling you something true about the solver.", None),
 ("", None),
 ("!! THE SCALE WARNING - READ THIS BEFORE CONCLUDING ANYTHING", F_H1),
 ("Small models are TIGHT; large ones are LOOSE. This toy has 3 crops and 5 plantings. The", None),
 ("real farm has 7 crops and 220 acre-plantings. Which constraints BIND, and which crops are", None),
 ("worth planting, change with scale, and this sheet will get that wrong.", None),
 ("", None),
 ("So: this toy teaches MECHANISM - how storage works, what it measures, what it takes away.", None),
 ("It does NOT tell you which constraint is actually binding on your farm. The app already", None),
 ("answers that (capacity_series returns used-vs-cap per week for every pool). Any conclusion", None),
 ("drawn here gets re-measured against the real farm before it becomes a plan item.", None),
 ("", None),
 ("CHECK IT WORKED - the numbers you should see on opening", F_H1),
 ("With the sheet as shipped (every planting 6 acres, sell rate 1,000,000, shrink 0%):", None),
 ("     Peak COLD store         66.6 m3  in week 41      (cooler holds 275)", F_BOLD),
 ("     Peak AMBIENT store      61.0 m3  in week 30      (barn holds 150)", F_BOLD),
 ("     Stock carried out       0 lb in every week, for all three crops", F_BOLD),
 ("     Spinach 51,000 lb sold | Kale 81,600 lb | Garlic 51,000 lb", None),
 ("If those match, the sheet is calculating correctly. If they do not, tell me - do not work", None),
 ("around it. A toy you cannot trust is worse than no toy.", None),
 ("", None),
 ("Note the pair in bold: closing stock is ZERO all year, and the cooler is still 66.6 m3 full.", None),
 ("", None),
 ("THREE THINGS TO TRY", F_H1),
 ("Each one changes ONE input. Change it, look, then put it back.", None),
 ("", None),
 ("EXPERIMENT 1 - overflow the stores.  Set all five acre cells (B6:B10) to 25.", F_BOLD),
 ("   Expect: peak COLD 277.6 m3 in week 41, red for 1 week. Peak AMBIENT 254.0 m3 in week 30,", None),
 ("   red for 1 week. The app would refuse this plan outright - not warn about it, refuse it.", None),
 ("   Note it takes ~24.8 acres per planting to break the cooler, and the break is in ONE week.", None),
 ("   That is what a capacity constraint feels like: nothing, nothing, nothing, infeasible.", None),
 ("", None),
 ("EXPERIMENT 2 - hold stock, and watch what does NOT happen.  Put acres back to 6.", F_BOLD),
 ("   Set Garlic's 'Sell rate lb/wk' (H17) to 2000.", F_BOLD),
 ("   Expect: Garlic now sells 46,000 lb instead of 51,000, carries 5,000 lb into next year,", None),
 ("   and 'Over shelf life?' reaches 19,000 lb.", None),
 ("   BUT peak AMBIENT is still 61.0 m3 in week 30 - EXACTLY as before.", F_BOLD),
 ("   Holding stock for months did not raise the peak by a single cubic metre, because the peak", None),
 ("   is set by the week the harvest lands, not by how long you keep it. This is the same point", None),
 ("   as the opening one, arriving from the other direction: size the barn for harvest day.", None),
 ("", None),
 ("EXPERIMENT 3 - make holding cost something.  Leave sell rate at 2000.", F_BOLD),
 ("   Set Garlic's 'Shrink %/wk' (G17) to 2%.", F_BOLD),
 ("   Expect: Garlic sells 41,474 lb (down 4,526), loses 9,526 lb to shrink, and carries NOTHING", None),
 ("   into next year - the shrink ate the leftovers. 'Over shelf life' falls to 10,047 lb.", None),
 ("   Without shrink, a long shelf life is pure upside and the model would hoard. Shrink is what", None),
 ("   makes storing a real decision with a real cost.", None),
 ("", None),
 ("WHERE THE NUMBERS COME FROM", F_H1),
 ("Every crop constant and every planting calendar on the Model sheet is READ FROM THE APP by", None),
 ("toy_models/build_toy_v1.py - the same succession generator the solver uses, and the same", None),
 ("golden payload the tests pin. Nothing is retyped. See the 'Crop data' sheet for the values", None),
 ("and their sources. Re-run the generator after any crop-data change.", None),
]
for i, (text, font) in enumerate(readme, start=1):
    c = rd.cell(row=i, column=1, value=text)
    if font: c.font = font

# ===========================================================================
# SHEET 2 — MODEL
# ===========================================================================
ws = wb.create_sheet("Model")
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 13
for w in range(1, 53):
    ws.column_dimensions[wl(w)].width = 8.5

def put(r, c, v, font=None, fill=None, fmt=None, align=None, border=False):
    cell = ws.cell(row=r, column=c, value=v)
    if font: cell.font = font
    if fill: cell.fill = fill
    if fmt:  cell.number_format = fmt
    if align: cell.alignment = Alignment(horizontal=align)
    if border: cell.border = BOX
    return cell

r = 1
put(r, 1, "FARM TOY MODEL — v1: the spine + storage"); ws.cell(row=r, column=1).font = F_TITLE
r += 1
put(r, 1, "Type in the YELLOW cells only. Everything else is a formula. Read the 'READ ME' sheet first.", font=F_NOTE)
r += 2

# ---- INPUT: acres per succession ----
put(r, 1, "1 — WHAT YOU PLANT   (acres)", font=F_H1); r += 1
hdr = ["Succession", "Acres", "Crop", "Harvest weeks", "lb/harvest wk at 1 acre"]
for j, h in enumerate(hdr):
    put(r, 1 + j, h, font=F_BOLD, fill=HDR_FILL, border=True)
r += 1
ROW_SUCC = r
succ_meta = []
for s in succs:
    crop = crops[s["crop_name"]] if "crop_name" in s else crops[CROP_NAMES[s["crop"]]]
    nharv = max(1, len(harvest_weeks_of(s, 52)))
    put(r, 1, s["name"], border=True)
    put(r, 2, 6.0, fill=IN_FILL, fmt="0.0", border=True)          # <- INPUT
    put(r, 3, crop["name"], border=True)
    put(r, 4, ", ".join("W%d" % w for w in harvest_weeks_of(s, 52)), font=F_SMALL, border=True)
    put(r, 5, round(crop["yield_per_acre"] * STF / nharv, 1), fmt="#,##0", border=True)
    succ_meta.append({"row": r, "s": s, "crop": crop["name"], "nharv": nharv})
    r += 1
put(r, 1, "Total acres planted", font=F_BOLD)
put(r, 2, f"=SUM(B{ROW_SUCC}:B{r-1})", font=F_BOLD, fmt="0.0")
put(r, 3, f"of {TOTAL_LAND} acres of ground", font=F_NOTE)
ROW_TOTACRES = r
r += 2

# ---- INPUT: crop settings ----
put(r, 1, "2 — HOW EACH CROP BEHAVES", font=F_H1); r += 1
cols = ["Crop", "Yield lb/acre", "Sell-through", "m3 per lb", "Store", "Shelf life wk",
        "Shrink %/wk", "Sell rate lb/wk"]
for j, h in enumerate(cols):
    put(r, 1 + j, h, font=F_BOLD, fill=HDR_FILL, border=True)
r += 1
ROW_CROP = r
crop_row = {}
for name in CROP_NAMES:
    c = crops[name]
    pool = "Cold" if c.get("cold_storage", True) else "Ambient"
    put(r, 1, name, border=True)
    put(r, 2, c["yield_per_acre"], fmt="#,##0", border=True)
    put(r, 3, STF, fill=IN_FILL, fmt="0.00", border=True)
    put(r, 4, c["crop_size_m3_per_lb"], fmt="0.0000000", border=True)
    put(r, 5, pool, border=True, fill=POOL_C if pool == "Cold" else POOL_A)
    put(r, 6, c.get("shelf_life_weeks", c.get("sales_spread_weeks", 1)), fill=IN_FILL, border=True)
    put(r, 7, 0.0, fill=IN_FILL, fmt="0.0%", border=True)
    put(r, 8, 1000000, fill=IN_FILL, fmt="#,##0", border=True)
    crop_row[name] = r
    r += 1
put(r, 1, "Shrink is weight lost in store each week. Sell rate stands in for the buyer (layer 2 makes it real).", font=F_NOTE)
r += 2

# ---- INPUT: storage caps ----
put(r, 1, "3 — WHAT YOU CAN STORE", font=F_H1); r += 1
put(r, 1, "Cold store (cooler) capacity", border=True)
put(r, 2, COLD_CAP, fill=IN_FILL, fmt="#,##0", border=True); ROW_COLDCAP = r
put(r, 3, "m3   — Spinach, Kale", font=F_NOTE); r += 1
put(r, 1, "Ambient store (barn) capacity", border=True)
put(r, 2, AMB_CAP, fill=IN_FILL, fmt="#,##0", border=True); ROW_AMBCAP = r
put(r, 3, "m3   — Garlic (cured; the cooler would damage it)", font=F_NOTE)
r += 2

# ---- week header ----
put(r, 1, "4 — THE YEAR, WEEK BY WEEK", font=F_H1); r += 1
ROW_WEEKNUM = r
put(r, 1, "Week", font=F_BOLD, fill=HDR_FILL)
put(r, 2, "", fill=HDR_FILL)
for w in range(1, 53):
    put(r, wcol(w), w, font=F_BOLD, fill=HDR_FILL, align="center")
ws.freeze_panes = ws.cell(row=r + 1, column=3)
r += 1

# ---- per-succession phase + harvest ----
put(r, 1, "PLANTING CALENDAR", font=F_BOLD); r += 1
PHASE_TXT = {0: "P", 1: "G", 2: "H"}
succ_harv_row = {}
for m in succ_meta:
    s = m["s"]
    phase_by_week = {}
    for wk, ph in zip(s["weeks"], s["phases"]):
        ww = ((wk - 1) % 52) + 1
        phase_by_week[ww] = ph
    put(r, 1, s["name"], font=F_SMALL)
    put(r, 2, "P/G/H", font=F_NOTE)
    for w in range(1, 53):
        if w in phase_by_week:
            put(r, wcol(w), PHASE_TXT[phase_by_week[w]], align="center", font=F_SMALL,
                fill=CALC_FILL)
    phrow = r
    r += 1
    put(r, 1, "   lbs harvested (sellable)", font=F_SMALL)
    put(r, 2, "lb", font=F_NOTE)
    for w in range(1, 53):
        col = wl(w)
        put(r, wcol(w),
            f'=IF({col}{phrow}="H",$B${m["row"]}*$B${crop_row[m["crop"]]}*$C${crop_row[m["crop"]]}/{m["nharv"]},0)',
            fmt="#,##0")
    succ_harv_row.setdefault(m["crop"], []).append(r)
    r += 1
put(r, 1, "P = planting   G = growing   H = harvesting.  Only sellable pounds are shown (yield x sell-through).", font=F_NOTE)
r += 2

# ---- per-crop inventory + storage ----
crop_block = {}
for name in CROP_NAMES:
    cr = crop_row[name]
    pool = "Cold" if crops[name].get("cold_storage", True) else "Ambient"
    put(r, 1, f"{name.upper()}  —  {pool.lower()} store", font=F_H1)
    put(r, 2, "", fill=POOL_C if pool == "Cold" else POOL_A)
    r += 1

    put(r, 1, "Harvested this week", font=F_SMALL); put(r, 2, "lb", font=F_NOTE)
    for w in range(1, 53):
        parts = "+".join(f"{wl(w)}{hr}" for hr in succ_harv_row[name])
        put(r, wcol(w), f"={parts}", fmt="#,##0")
    R_HARV = r; r += 1

    put(r, 1, "Stock carried in", font=F_SMALL); put(r, 2, "lb", font=F_NOTE)
    for w in range(1, 53):
        put(r, wcol(w), 0 if w == 1 else f"={wl(w-1)}{r+4}", fmt="#,##0")
    R_OPEN = r; r += 1

    put(r, 1, "   lost to shrink", font=F_SMALL); put(r, 2, "lb", font=F_NOTE)
    for w in range(1, 53):
        put(r, wcol(w), f"={wl(w)}{R_OPEN}*$G${cr}", fmt="#,##0")
    R_SHRINK = r; r += 1

    put(r, 1, "Available to sell", font=F_SMALL); put(r, 2, "lb", font=F_NOTE)
    for w in range(1, 53):
        put(r, wcol(w), f"={wl(w)}{R_OPEN}-{wl(w)}{R_SHRINK}+{wl(w)}{R_HARV}", fmt="#,##0")
    R_AVAIL = r; r += 1

    put(r, 1, "Sold", font=F_SMALL); put(r, 2, "lb", font=F_NOTE)
    for w in range(1, 53):
        put(r, wcol(w), f"=MIN({wl(w)}{R_AVAIL},$H${cr})", fmt="#,##0")
    R_SOLD = r; r += 1

    put(r, 1, "Stock carried out", font=F_BOLD); put(r, 2, "lb", font=F_NOTE)
    for w in range(1, 53):
        put(r, wcol(w), f"={wl(w)}{R_AVAIL}-{wl(w)}{R_SOLD}", fmt="#,##0")
    R_CLOSE = r; r += 1

    put(r, 1, "ROOM OCCUPANCY", font=F_BOLD)
    put(r, 2, "m3", font=F_NOTE)
    for w in range(1, 53):
        put(r, wcol(w), f"=({wl(w)}{R_OPEN}+{wl(w)}{R_HARV})*$D${cr}", fmt="0.0",
            fill=POOL_C if pool == "Cold" else POOL_A)
    R_ROOM = r; r += 1
    put(r, 1, "   = carried in + ALL of this week's harvest, before any of it ships", font=F_NOTE)
    r += 1

    put(r, 1, "Over shelf life?", font=F_SMALL); put(r, 2, "lb", font=F_NOTE)
    for w in range(1, 53):
        put(r, wcol(w),
            f'=MAX(0,{wl(w)}{R_CLOSE}-SUMIFS(${wl(1)}${R_HARV}:${LAST}${R_HARV},'
            f'${wl(1)}${ROW_WEEKNUM}:${LAST}${ROW_WEEKNUM},">="&{wl(w)}${ROW_WEEKNUM}-$F${cr}+2,'
            f'${wl(1)}${ROW_WEEKNUM}:${LAST}${ROW_WEEKNUM},"<="&{wl(w)}${ROW_WEEKNUM}))',
            fmt="#,##0")
    R_OLD = r; r += 1
    put(r, 1, "   stock older than its shelf life — a WARNING here, a hard constraint in the app", font=F_NOTE)
    r += 2
    crop_block[name] = dict(harv=R_HARV, sold=R_SOLD, close=R_CLOSE, room=R_ROOM,
                            shrink=R_SHRINK, old=R_OLD, pool=pool)

# ---- pools ----
put(r, 1, "5 — THE TWO STORES", font=F_H1); r += 1
pool_rows = {}
for pool, cap_row, fill in (("Cold", ROW_COLDCAP, POOL_C), ("Ambient", ROW_AMBCAP, POOL_A)):
    members = [n for n in CROP_NAMES if crop_block[n]["pool"] == pool]
    put(r, 1, f"{pool.upper()} POOL used", font=F_BOLD, fill=fill)
    put(r, 2, "m3", font=F_NOTE)
    for w in range(1, 53):
        parts = "+".join(f"{wl(w)}{crop_block[n]['room']}" for n in members)
        put(r, wcol(w), f"={parts}", fmt="0.0", fill=fill)
    used = r; r += 1
    put(r, 1, f"{pool.upper()} POOL headroom", font=F_BOLD)
    put(r, 2, "m3", font=F_NOTE)
    for w in range(1, 53):
        put(r, wcol(w), f"=$B${cap_row}-{wl(w)}{used}", fmt="0.0")
    head = r; r += 1
    put(r, 1, f"   {pool} store holds {', '.join(members)}. NEGATIVE = overflowed = the app calls this infeasible.",
        font=F_NOTE)
    r += 2
    ws.conditional_formatting.add(
        f"{wl(1)}{head}:{LAST}{head}",
        CellIsRule(operator="lessThan", formula=["0"],
                   fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(
        f"{wl(1)}{used}:{LAST}{used}",
        CellIsRule(operator="greaterThan", formula=[f"$B${cap_row}"],
                   fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006", bold=True)))
    pool_rows[pool] = dict(used=used, head=head, cap=cap_row)

# ---- summary ----
put(r, 1, "6 — WHAT THIS PLAN DOES", font=F_H1); r += 1
def summary(label, formula, fmt="#,##0", note="", bold=False):
    global r
    put(r, 1, label, font=F_BOLD if bold else None)
    put(r, 2, formula, fmt=fmt, font=F_BOLD if bold else None)
    if note: put(r, 3, note, font=F_NOTE)
    r += 1

summary("Peak COLD store used", f"=MAX({wl(1)}{pool_rows['Cold']['used']}:{LAST}{pool_rows['Cold']['used']})",
        "0.0", f"m3   against a {COLD_CAP} m3 cooler", bold=True)
summary("Weeks the cooler overflows",
        f"=COUNTIF({wl(1)}{pool_rows['Cold']['head']}:{LAST}{pool_rows['Cold']['head']},\"<0\")",
        "0", "weeks out of 52", bold=True)
summary("Peak AMBIENT store used", f"=MAX({wl(1)}{pool_rows['Ambient']['used']}:{LAST}{pool_rows['Ambient']['used']})",
        "0.0", f"m3   against a {AMB_CAP} m3 barn", bold=True)
summary("Weeks the barn overflows",
        f"=COUNTIF({wl(1)}{pool_rows['Ambient']['head']}:{LAST}{pool_rows['Ambient']['head']},\"<0\")",
        "0", "weeks out of 52", bold=True)
r += 1
for name in CROP_NAMES:
    b = crop_block[name]
    summary(f"{name} — harvested", f"=SUM({wl(1)}{b['harv']}:{LAST}{b['harv']})", "#,##0", "lb sellable")
    summary(f"{name} — sold",      f"=SUM({wl(1)}{b['sold']}:{LAST}{b['sold']})", "#,##0", "lb")
    summary(f"{name} — lost to shrink", f"=SUM({wl(1)}{b['shrink']}:{LAST}{b['shrink']})", "#,##0", "lb")
    summary(f"{name} — left at year end", f"={LAST}{b['close']}", "#,##0",
            "lb  (the app anchors stock at zero on 1 Jan, so this is stranded)")
    summary(f"{name} — over shelf life at worst", f"=MAX({wl(1)}{b['old']}:{LAST}{b['old']})", "#,##0",
            "lb  (WARNING only in this sheet)")
    r += 1

# ===========================================================================
# SHEET 3 — CROP DATA
# ===========================================================================
cd = wb.create_sheet("Crop data")
cd.column_dimensions["A"].width = 34
for col in "BCD":
    cd.column_dimensions[col].width = 17
rows = [("Crop constants — READ FROM THE APP, not retyped", F_TITLE)]
cd.cell(row=1, column=1, value=rows[0][0]).font = F_TITLE
cd.cell(row=2, column=1,
        value="Source: tests/golden/default_farm_payload.json (the payload the browser sends) and "
              "farm_v2.utils.calendar.generate_successions (the app's own calendar).").font = F_NOTE
cd.cell(row=3, column=1, value="Regenerate with:  python3 toy_models/build_toy_v1.py").font = F_NOTE
hdrs = ["Field", "Spinach", "Kale", "Garlic", "What it means"]
for j, h in enumerate(hdrs):
    c = cd.cell(row=5, column=1 + j, value=h); c.font = F_BOLD; c.fill = HDR_FILL
fields = [
 ("yield_per_acre", "lb of crop per acre, before sell-through"),
 ("crop_size_m3_per_lb", "how much store space one pound takes"),
 ("cold_storage", "TRUE = cooler, FALSE = ambient barn"),
 ("shelf_life_weeks", "weeks it can be SOLD IN, counting its harvest week"),
 ("days_to_maturity", "plant to FIRST harvest (not total field time)"),
 ("planting_weeks", "how long the planting phase lasts"),
 ("harvest_weeks", "how many weeks it is picked over"),
 ("overwinters", "may its cycle run past week 52"),
 ("seasons", "the PLANTING window(s) — any start week inside is legal"),
]
rr = 6
for key, meaning in fields:
    cd.cell(row=rr, column=1, value=key).font = F_BOLD
    for j, name in enumerate(CROP_NAMES):
        v = crops[name].get(key, crops[name].get("sales_spread_weeks") if key == "shelf_life_weeks" else "")
        if key == "seasons":
            v = "; ".join(f"{s['start']}->{s['end']}" for s in crops[name]["seasons"])
        cd.cell(row=rr, column=2 + j, value=str(v) if not isinstance(v, (int, float)) else v)
    cd.cell(row=rr, column=5, value=meaning).font = F_SMALL
    rr += 1
rr += 1
cd.cell(row=rr, column=1, value="Farm-level").font = F_H1; rr += 1
for k, v, note in (("total_land", TOTAL_LAND, "acres of ground"),
                   ("sell_through_fraction", STF, "share of harvest that becomes sellable product"),
                   ("harvest_storage_cap", COLD_CAP, "m3 of refrigerated space"),
                   ("ambient_storage_cap", AMB_CAP, "m3 of barn space")):
    cd.cell(row=rr, column=1, value=k).font = F_BOLD
    cd.cell(row=rr, column=2, value=v)
    cd.cell(row=rr, column=5, value=note).font = F_SMALL
    rr += 1
rr += 2
cd.cell(row=rr, column=1, value="The 5 plantings on the Model sheet").font = F_H1; rr += 1
cd.cell(row=rr, column=1, value="chosen from the 68 the app's generator offers for these three crops").font = F_NOTE
rr += 1
for j, h in enumerate(["Succession", "Crop", "In the ground", "Harvest weeks"]):
    c = cd.cell(row=rr, column=1 + j, value=h); c.font = F_BOLD; c.fill = HDR_FILL
rr += 1
for m in succ_meta:
    s = m["s"]
    cd.cell(row=rr, column=1, value=s["name"])
    cd.cell(row=rr, column=2, value=m["crop"])
    cd.cell(row=rr, column=3, value=f"W{s['weeks'][0]} to W{((s['weeks'][-1]-1)%52)+1} ({len(s['weeks'])} wk)")
    cd.cell(row=rr, column=4, value=", ".join("W%d" % w for w in harvest_weeks_of(s, 52)))
    rr += 1

out = REPO / "toy_models" / "farm_toy_v1_spine_storage.xlsx"
wb.save(out)
print(f"wrote {out}")
print(f"  crops: {', '.join(CROP_NAMES)}")
print(f"  successions: {', '.join(PICKED)}")
print(f"  cold cap {COLD_CAP} m3 | ambient cap {AMB_CAP} m3 | sell-through {STF}")

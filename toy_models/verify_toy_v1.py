"""
verify_toy_v1.py — EVALUATE the toy workbook and prove it works.

Run:  .venv/bin/python toy_models/verify_toy_v1.py
      (needs `formulas`; `tests/test_toy_model.py` runs this as a subprocess)

WHY THIS EXISTS. openpyxl reads formulas as TEXT. Every static check — "is this a formula",
"does it reference the right row" — can pass over a sheet that computes nonsense. The grower
opened v1 and within a minute found a column of frozen numbers sitting beside the live ones
they described, then said: if there is an error right at the beginning, there are more. That
is the correct inference, and the only way to answer it is to CALCULATE the sheet.

Three things are checked, in increasing strength:

  A. VALUES — every computed row equals an independent Python simulator built from the app's
     own crop data. Catches a formula that is well-formed and wrong.
  B. PROPAGATION — perturb each input, recalculate, and assert the outputs actually move.
     Catches an input wired to nothing, which is what E6:E10 was.
  C. NO FROZEN DERIVED VALUES — every number on the Model sheet is a formula, a declared
     input, or one of three whitelisted constants (the week strips, the week-1 stock anchor).

⚠️ THE TRAP THAT COST AN HOUR HERE, WRITTEN DOWN SO IT IS NOT REPEATED.
`formulas` keys cells as "'[filename.xlsx]SHEET'!A1" and is CASE-SENSITIVE on the filename.
An override posted with an uppercased name does not fail — it creates a NEW orphan node, the
real cell keeps its old value, and reading the key back returns the value just injected. So
the check confirms itself and every propagation test reports a pass over a sheet nothing was
done to. Both the case and a landed-assertion are enforced in `calc()` below.
"""
import json
import re
import sys
import pathlib
import warnings

warnings.filterwarnings("ignore")
REPO = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import formulas
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from farm_v2.utils.calendar import generate_successions, harvest_weeks_of

WB = REPO / "toy_models" / "farm_toy_v1_spine_storage.xlsx"
PFX = f"'[{WB.name}]MODEL'!"          # exact case — see the trap note above
CROPS = ("Spinach", "Kale", "Garlic")
PICKED = [("Spinach W10", "Spinach"), ("Spinach W34", "Spinach"),
          ("Kale W12", "Kale"), ("Kale W32", "Kale"), ("Garlic W47", "Garlic")]
CELL_RE = re.compile(rf"'\[{re.escape(WB.name)}\]MODEL'!([A-Z]+\d+)$")

ws = load_workbook(WB)["Model"]
failures = []


def report(ok, msg):
    print(("  OK   " if ok else "  FAIL ") + msg)
    if not ok:
        failures.append(msg)


def labels():
    out = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str):
            out.setdefault(v.strip(), []).append(r)
    return out


def crop_blocks():
    out = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and "—" in v and v.endswith("store"):
            seq = {}
            for rr in range(r + 1, r + 14):
                lv = ws.cell(row=rr, column=1).value
                if isinstance(lv, str):
                    seq[lv.strip()] = rr
            out[v.split(" ")[0]] = seq
    return out


LBL, BLK = labels(), crop_blocks()
col = lambda w: get_column_letter(w + 2)
xl = formulas.ExcelModel().loads(str(WB)).finish()


def calc(overrides=None):
    sol = (xl.calculate(inputs={PFX + k: v for k, v in overrides.items()})
           if overrides else xl.calculate())
    d = {}
    for k, v in sol.items():
        m = CELL_RE.match(k)
        if m:
            try:
                d[m.group(1)] = v.value[0, 0]
            except Exception:
                pass
    for k, v in (overrides or {}).items():
        got = d.get(k)
        assert got is not None and abs(float(got) - float(v)) < 1e-9, (
            f"OVERRIDE DID NOT LAND: {k} reads {got!r}, wanted {v}. Check the key case — an "
            f"uppercased filename silently creates an orphan node instead of overriding.")
    return d


def row(d, r):
    return [d.get(f"{col(w)}{r}") for w in range(1, 53)]


# --------------------------------------------------------------------------- simulator
payload = json.load(open(REPO / "tests" / "golden" / "default_farm_payload.json"))
crops = {c["name"]: c for c in payload["crops"] if c["name"] in CROPS}
gen = {s["name"]: s for s in generate_successions(list(crops.values()), 2024)}


def simulate(acres, stf, sell, shrink):
    harv = {n: [0.0] * 53 for n in crops}
    land = [0.0] * 53
    for pn, cn in PICKED:
        s = gen[pn]
        hw = harvest_weeks_of(s, 52)
        for w in hw:
            harv[cn][w] += acres[pn] * crops[cn]["yield_per_acre"] * stf[cn] / max(1, len(hw))
        for wk in s["weeks"]:
            land[((wk - 1) % 52) + 1] += acres[pn]
    close = {n: [0.0] * 53 for n in crops}
    room = {n: [0.0] * 53 for n in crops}
    for n, c in crops.items():
        for w in range(1, 53):
            opening = close[n][w - 1] if w > 1 else 0.0
            avail = opening - opening * shrink[n] + harv[n][w]
            close[n][w] = avail - min(avail, sell[n])
            room[n][w] = (opening + harv[n][w]) * c["crop_size_m3_per_lb"]
    return dict(harv=harv, land=land, room=room, close=close,
                cold=[room["Spinach"][w] + room["Kale"][w] for w in range(1, 53)],
                amb=[room["Garlic"][w] for w in range(1, 53)])


BASE = dict(acres={pn: 6.0 for pn, _ in PICKED}, stf={n: 0.85 for n in crops},
            sell={n: 1e6 for n in crops}, shrink={n: 0.0 for n in crops})


def compare(tag, got, want, tol=1e-6):
    want = want[1:] if len(want) == 53 else want
    bad = [(w + 1, g, x) for w, (g, x) in enumerate(zip(got, want))
           if g is None or abs(float(g) - x) > tol]
    report(not bad, tag + ("" if not bad else
           f" — first mismatch wk{bad[0][0]}: sheet {bad[0][1]}, simulator {bad[0][2]:.4f} "
           f"({len(bad)} weeks differ)"))


def check_values(tag, overrides, params):
    d = calc(overrides)
    s = simulate(**params)
    for n in CROPS:
        b = BLK[n.upper()]
        compare(f"[{tag}] {n} harvested",         row(d, b["Harvested this week"]), s["harv"][n])
        compare(f"[{tag}] {n} stock carried out", row(d, b["Stock carried out"]),   s["close"][n])
        compare(f"[{tag}] {n} room occupancy",    row(d, b["ROOM OCCUPANCY"]),      s["room"][n], 1e-4)
    compare(f"[{tag}] LAND in the ground", row(d, LBL["LAND in the ground"][0]), s["land"])
    compare(f"[{tag}] COLD STORE used",    row(d, LBL["COLD STORE"][0]),    [0] + s["cold"], 1e-4)
    compare(f"[{tag}] AMBIENT STORE used", row(d, LBL["AMBIENT STORE"][0]), [0] + s["amb"],  1e-4)
    return d


print("A. VALUES — the sheet against an independent simulator")
print("   TWO configurations, and the second is not optional. At the shipped defaults shrink is")
print("   0 and every sell rate is effectively infinite, so a whole class of sign and ordering")
print("   errors is INERT: `opening - 0 + harvest` and `opening + 0 + harvest` agree exactly.")
print("   Measured — flipping the shrink sign in the generator passed this section until the")
print("   stressed configuration was added.")
d0 = check_values("defaults", None, BASE)

STRESS_OV = {"B6": 10.0, "B10": 8.0, "C16": 0.50, "H18": 2000.0, "G18": 0.02,
             "H16": 9000.0, "G17": 0.05}
STRESS_P = dict(
    acres={**{pn: 6.0 for pn, _ in PICKED}, "Spinach W10": 10.0, "Garlic W47": 8.0},
    stf={"Spinach": 0.50, "Kale": 0.85, "Garlic": 0.85},
    sell={"Spinach": 9000.0, "Kale": 1e6, "Garlic": 2000.0},
    shrink={"Spinach": 0.0, "Kale": 0.05, "Garlic": 0.02})
check_values("stressed", STRESS_OV, STRESS_P)

print("\nB. PROPAGATION — perturb each input, does anything downstream move?")
peak = lambda d, lab, i=0: max(float(x) for x in row(d, LBL[lab][i]) if x is not None)
total = lambda d, c, r: sum(float(x) for x in row(d, BLK[c][r]) if x is not None)
worst = lambda d, i: min(float(x) for x in row(d, LBL["room left"][i]) if x is not None)
base_sp, base_cold = total(d0, "SPINACH", "Harvested this week"), peak(d0, "COLD STORE")

for tag, ov, pred in [
    ("acres         B6  6 -> 12",        {"B6": 12.0},
     lambda d: total(d, "SPINACH", "Harvested this week") > base_sp * 1.4),
    ("sell-through  C16 0.85 -> 0.50",   {"C16": 0.50},
     lambda d: total(d, "SPINACH", "Harvested this week") < base_sp * 0.7),
    ("yield         B16 5000 -> 10000",  {"B16": 10000},
     lambda d: total(d, "SPINACH", "Harvested this week") > base_sp * 1.8),
    ("m3 per lb     D16 x2",             {"D16": crops["Spinach"]["crop_size_m3_per_lb"] * 2},
     lambda d: peak(d, "COLD STORE") > base_cold * 1.2),
    ("sell rate     H18 -> 2000",        {"H18": 2000.0},
     lambda d: total(d, "GARLIC", "Stock carried out") > 0),
    ("shrink        G18 -> 2%",          {"H18": 2000.0, "G18": 0.02},
     lambda d: total(d, "GARLIC", "lost to shrink") > 0),
    ("cold cap      B22 275 -> 50",      {"B22": 50.0},  lambda d: worst(d, 1) < 0),
    ("ambient cap   B23 150 -> 20",      {"B23": 20.0},  lambda d: worst(d, 2) < 0),
    ("land          B12 80 -> 10",       {"B12": 10.0},  lambda d: worst(d, 0) < 0),
    ("shelf life    F18 16 -> 2",        {"H18": 2000.0, "F18": 2},
     lambda d: total(d, "GARLIC", "Over shelf life?") > 0),
]:
    try:
        ok = pred(calc(ov))
    except AssertionError as e:
        ok, tag = False, f"{tag}  [{e}]"
    except Exception as e:
        ok, tag = False, f"{tag}  [{type(e).__name__}: {e}]"
    report(ok, tag)

print("\nC. NO FROZEN DERIVED VALUES")
week_rows = {r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Week"}
anchor = {r for r in range(1, ws.max_row + 1)
          if str(ws.cell(row=r, column=1).value).strip() == "Stock carried in"}
frozen = []
for line in ws.iter_rows():
    for c in line:
        if not isinstance(c.value, (int, float)) or isinstance(c.value, bool):
            continue
        yellow = bool(c.fill and c.fill.fgColor and c.fill.fgColor.rgb
                      and "FFF2CC" in str(c.fill.fgColor.rgb))
        if yellow or c.row in week_rows or (c.row in anchor and c.column == 3):
            continue
        frozen.append(f"{c.coordinate}={c.value} (row {ws.cell(row=c.row, column=1).value!r})")
report(not frozen, f"every number is a formula or a declared input"
                   + ("" if not frozen else " — FROZEN: " + "; ".join(frozen[:6])))

print("\n" + ("ALL CHECKS PASSED" if not failures else f"{len(failures)} FAILURE(S)"))
sys.exit(1 if failures else 0)
